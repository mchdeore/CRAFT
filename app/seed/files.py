"""Write synthetic source formats with fixed properties and package bytes."""

import csv
import html
import textwrap
from io import BytesIO, StringIO
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


def stable_package(content: bytes) -> bytes:
    with ZipFile(BytesIO(content)) as package:
        entries = {name: package.read(name) for name in package.namelist()}
    if "docProps/core.xml" in entries:
        core = ElementTree.fromstring(entries["docProps/core.xml"])
        for node in core:
            name = node.tag.rsplit("}", 1)[-1]
            if name in {"created", "modified"}:
                node.text = "2026-01-01T00:00:00Z"
            if name in {"creator", "lastModifiedBy"}:
                node.text = ""
        entries["docProps/core.xml"] = ElementTree.tostring(
            core, encoding="utf-8", xml_declaration=True
        )
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as package:
        for name, data in sorted(entries.items()):
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            package.writestr(info, data)
    return output.getvalue()


def document(title: str, paragraphs: list[str], format: str) -> bytes:
    if format in {"md", "txt"}:
        heading = "# " if format == "md" else ""
        return (heading + title + "\n\n" + "\n\n".join(paragraphs) + "\n").encode()
    if format == "html":
        body = "".join("<p>" + html.escape(part) + "</p>" for part in paragraphs)
        return (
            "<!doctype html><html><meta charset='utf-8'><body><h1>"
            + html.escape(title)
            + "</h1>"
            + body
            + "</body></html>"
        ).encode()
    if format == "docx":
        from docx import Document

        source = Document()
        source.add_heading("1 " + title, level=1)
        for part in paragraphs[:4]:
            source.add_paragraph(part)
        source.add_heading("2 Verification and assumptions", level=1)
        for part in paragraphs[4:]:
            source.add_paragraph(part)
        source.core_properties.title = title
        source.core_properties.author = ""
        source.core_properties.last_modified_by = ""
        output = BytesIO()
        source.save(output)
        return stable_package(output.getvalue())
    if format == "pptx":
        from pptx import Presentation
        from pptx.shapes.autoshape import Shape

        source = Presentation()
        for number in range(3):
            slide = source.slides.add_slide(source.slide_layouts[1])
            heading = slide.shapes.title
            assert heading is not None
            heading.text = title + f" / {number + 1}"
            body = slide.placeholders[1]
            assert isinstance(body, Shape)
            body.text = paragraphs[min(number, len(paragraphs) - 1)]
            notes = slide.notes_slide.notes_text_frame
            assert notes is not None
            notes.text = "\n".join(paragraphs[number::3])
        source.core_properties.title = title
        source.core_properties.author = ""
        source.core_properties.last_modified_by = ""
        output = BytesIO()
        source.save(output)
        return stable_package(output.getvalue())
    if format == "pdf":
        return pdf(title, paragraphs)
    raise ValueError("unsupported synthetic format")


def pdf(title: str, paragraphs: list[str]) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    source = PdfWriter()
    source.add_metadata(
        {
            "/Title": title,
            "/CreationDate": "D:20260101000000Z",
            "/ModDate": "D:20260101000000Z",
            "/Producer": "",
            "/Creator": "",
        }
    )
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    lines = [title, "SIMULATED / UNCLASSIFIED", ""]
    for part in paragraphs:
        lines.extend(textwrap.wrap(part, width=88) + [""])
    for start in range(0, len(lines), 45):
        page = source.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        commands = []
        for number, line in enumerate(lines[start : start + 45]):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"BT /F1 10 Tf 1 0 0 1 40 {750 - number * 15} Tm ({escaped}) Tj ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("cp1252"))
        page.replace_contents(stream.flate_encode())
    output = BytesIO()
    source.write(output)
    return output.getvalue()


def spreadsheet(
    title: str, date: str, headers: list[str], rows: list[list], *, copy: bool = False
) -> bytes:
    from openpyxl import Workbook

    source = Workbook()
    sheet = source.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append([title, "SIMULATED / UNCLASSIFIED"])
    sheet.append(["Valid from", date])
    sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    source.properties.title = title + (" / duplicate metadata" if copy else "")
    source.properties.creator = ""
    source.properties.lastModifiedBy = ""
    output = BytesIO()
    source.save(output)
    return stable_package(output.getvalue())


def comma_table(title: str, seed: int) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Simulated cost item", "Estimate (credits)", "Reserve (%)", "Scenario"])
    for number, item in enumerate(["Qualification", "Integration", "Operations"], 1):
        writer.writerow([title + " / " + item, 10000 + seed * number * 137, 10 + number, seed])
    return output.getvalue().encode()
