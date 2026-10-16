"""Exact source support, bounded inspection, identity, confidence gates and metadata-only logs."""

import asyncio
import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from functools import wraps

import msgspec
import pytest
from agent_framework import ChatResponse, Message
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.agent.fake import FakeChatClient
from app.infrastructure.database import metadata as m
from app.infrastructure.errors import ApiError, Forbidden, InvalidInput, NotFound, TooLarge
from app.infrastructure.logging import configure_logging
from app.infrastructure.shapes import RawFile
from app.library.handlers import index_provider
from app.summarize import handlers as summary
from app.summarize.service import supported_filing, supported_sentence


def run_async(function):
    @wraps(function)
    def run(*args, **kwargs):
        return asyncio.run(function(*args, **kwargs))

    return run


def selected_source(source):
    source["provider"].write_content(
        source["path"],
        b"# 1 Thermal\n\nThe radiator margin is 5 K.\n\n# 2 Power\n\nThe power budget is 120 W.\n",
    )
    index_provider(source["settings"])
    version = source["provider"].get_item(source["item"])["versions"][-1]
    return {"item": source["item"], "version": version["version"], "sha256": version["sha256"]}


def upload_file(kind, metadata=False, conflict=False):
    stream = io.BytesIO()
    if kind == "pdf":
        from fpdf import FPDF

        file = FPDF()
        file.add_page()
        file.set_font("Helvetica", size=12)
        file.cell(text="Simulated fact one.")
        file.add_page()
        file.cell(text="Simulated fact two.")
        data = bytes(file.output())
        media = "application/pdf"
    elif kind == "docx":
        from docx import Document

        file = Document()
        file.add_heading("Stage B" if conflict else "Simulated source", level=1)
        file.add_paragraph(
            "Mission M-SYN-1. Silo engineering. Stage B." if metadata else "Simulated fact one."
        )
        if conflict:
            table = file.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Field"
            table.rows[0].cells[1].text = "Value"
            row = table.add_row()
            row.cells[0].text = "Stage"
            row.cells[1].text = "C"
        else:
            file.add_heading("Second section", level=1)
            file.add_paragraph("Simulated fact two.")
        file.save(stream)
        data = stream.getvalue()
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from openpyxl import Workbook

        file = Workbook()
        assert file.active is not None
        file.active.title = "First"
        file.active.append(["Field", "Value"])
        file.active.append(["Simulated fact", "One"])
        second = file.create_sheet("Second")
        second.append(["Field", "Value"])
        second.append(["Simulated fact", "Two"])
        file.save(stream)
        data = stream.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return RawFile(content=data, file_name="source." + kind, media_type=media, simulated=True)


def counts(source):
    names = ["missions", "cube_values", "doc_versions", "chunks", "workspace_files", "drafts"]
    with source["owner"].connect() as connection:
        return {
            name: connection.execute(
                select(func.count()).select_from(m.metadata.tables[name])
            ).scalar_one()
            for name in names
        }


@run_async
async def test_summarize_selected_every_sentence_cites_block(cube_source):
    request = selected_source(cube_source)
    result = await summary.selected(cube_source["caller"], request)
    assert len(result.sentences) == 2
    for sentence in result.sentences:
        for citation in sentence.citations:
            correct = hashlib.sha256(citation.quote.encode()).hexdigest() == citation.block_version
            assert correct and citation.handle is not None
            assert (
                citation.handle.item == request["item"]
                and citation.handle.version == request["version"]
            )
            assert supported_sentence(sentence.text, citation.quote)


@run_async
@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx"])
async def test_summarize_upload_preserves_source_and_order(cube_source, kind):
    request = upload_file(kind)
    before = counts(cube_source)
    result = await summary.upload(cube_source["caller"], request)
    assert result.source.version == hashlib.sha256(request.content).hexdigest()
    assert result.source.item is None and result.source.provider_version is None
    blocks = [citation.block for sentence in result.sentences for citation in sentence.citations]
    assert blocks == sorted(blocks) and len(result.sentences) == 2
    assert counts(cube_source) == before
    for sentence in result.sentences:
        assert all(
            c.handle is None and supported_sentence(sentence.text, c.quote)
            for c in sentence.citations
        )


@run_async
async def test_summarize_metadata_suggestions_evidence_confidence(cube_source):
    before = counts(cube_source)
    result = await summary.upload(cube_source["caller"], upload_file("docx", metadata=True))
    expected = {"mission": "M-SYN-1", "silo": "engineering", "stage": "B"}
    for field, value in expected.items():
        suggestion = getattr(result, field)
        correct = suggestion.value == value
        assert correct and suggestion.confidence == 0 and suggestion.evidence
        assert all(
            supported_filing(field, value, c.quote + " " + c.heading_path)
            for c in suggestion.evidence
        )
    assert counts(cube_source) == before


@run_async
async def test_summarize_ambiguous_or_missing_metadata(cube_source):
    result = await summary.upload(cube_source["caller"], upload_file("docx", conflict=True))
    assert result.stage.value is None and result.stage.confidence == 0
    assert len(result.stage.evidence) >= 2
    assert result.mission.value is None and not result.mission.evidence


@run_async
async def test_summarize_unsupported_claim_is_refused(cube_source, monkeypatch):
    request = selected_source(cube_source)

    def invalid(self, messages):
        data = {
            "sentences": [{"text": "The radiator margin is 9 K.", "blocks": [1]}],
            "filings": [],
        }
        return ChatResponse(
            messages=[Message(role="assistant", contents=[json.dumps(data)])], model="fake"
        )

    monkeypatch.setattr(FakeChatClient, "reply", invalid)
    with pytest.raises(ApiError):
        await summary.selected(cube_source["caller"], request)
    assert not supported_sentence("The margin is 5 K.", "It is false that The margin is 5 K.")
    assert not supported_filing("stage", "B", "The budget has changed.")
    assert not supported_filing("stage", "", "Stage B.")


@run_async
async def test_summarize_low_confidence_no_draft(cube_source):
    request = selected_source(cube_source)
    before = counts(cube_source)
    result = await summary.selected(cube_source["caller"], request)
    assert result.confidence.score == 0 and result.confidence.level == "low"
    assert result.status == "insufficient_evidence" and result.sentences
    assert counts(cube_source) == before


def test_summarize_auth_origin_ownership_matrix(cube_source, monkeypatch):
    request = selected_source(cube_source)
    called = []
    original = FakeChatClient.reply

    def spy(self, messages):
        called.append(True)
        return original(self, messages)

    monkeypatch.setattr(FakeChatClient, "reply", spy)
    with TestClient(cube_source["app"], base_url="https://testserver") as http:
        headers = {"Origin": "https://testserver"}
        assert http.post("/api/summarize", json=request, headers=headers).status_code == 401
        assert cube_source["sign_in"](http).status_code == 200
        assert http.post("/api/summarize", json=request).status_code == 403
        assert (
            http.post(
                "/api/summarize", json=request, headers={"Origin": "https://other.invalid"}
            ).status_code
            == 403
        )
        missing = request | {"item": "itm_000000000000"}
        assert http.post("/api/summarize", json=missing, headers=headers).status_code == 404
        assert not called
        assert http.post("/api/summarize", json=request, headers=headers).status_code == 200
        with cube_source["owner"].begin() as connection:
            table = m.metadata.tables["auth_sessions"]
            connection.execute(
                table.update()
                .where(table.c.user_id == cube_source["caller"].user_id)
                .values(expires_at=datetime.now(UTC) - timedelta(seconds=1))
            )
        assert http.post("/api/summarize", json=request, headers=headers).status_code == 401


@run_async
async def test_summarize_limits_and_malformed_request(cube_source):
    request = selected_source(cube_source)
    with pytest.raises(NotFound):
        await summary.selected(cube_source["caller"], request | {"sha256": "0" * 64})
    with pytest.raises(InvalidInput):
        await summary.selected(cube_source["caller"], request | {"userId": "unknown"})
    with pytest.raises(InvalidInput):
        await summary.upload(
            cube_source["caller"],
            RawFile(
                content=b"", file_name="source.pdf", media_type="application/pdf", simulated=True
            ),
        )
    with pytest.raises(TooLarge):
        await summary.upload(
            cube_source["caller"],
            RawFile(
                content=b"x" * (25 * 1024 * 1024 + 1),
                file_name="source.pdf",
                media_type="application/pdf",
                simulated=True,
            ),
        )
    with pytest.raises(Forbidden):
        await summary.upload(
            cube_source["caller"], msgspec.structs.replace(upload_file("pdf"), simulated=False)
        )


@run_async
async def test_summarize_cross_feature_calls_enter_handlers(cube_source, monkeypatch):
    from app.library import handlers

    request = selected_source(cube_source)
    called = []
    original = handlers.inspect_file

    async def inspection(caller, request):
        called.append(True)
        return await original(caller, request)

    monkeypatch.setattr(handlers, "inspect_file", inspection)
    await summary.upload(cube_source["caller"], upload_file("pdf"))
    assert called
    assert (await summary.selected(cube_source["caller"], request)).source.item == request["item"]


@run_async
async def test_summarize_no_transaction_during_model_call(cube_source, monkeypatch):
    from sqlalchemy import event
    from sqlalchemy.orm import Session

    active = set()
    checked = []
    original = FakeChatClient.reply

    def started(session, *args):
        active.add(id(session))

    def ended(session, *args):
        active.discard(id(session))

    event.listen(Session, "after_begin", started)
    event.listen(Session, "after_transaction_end", ended)

    def spy(self, messages):
        checked.append(not active)
        return original(self, messages)

    monkeypatch.setattr(FakeChatClient, "reply", spy)
    try:
        await summary.selected(cube_source["caller"], selected_source(cube_source))
    finally:
        event.remove(Session, "after_begin", started)
        event.remove(Session, "after_transaction_end", ended)
    assert checked and all(checked)


@run_async
async def test_summarize_logging_chain_without_content(cube_source):
    from app.infrastructure.operations.storage import call_log

    request = selected_source(cube_source)
    stream = io.StringIO()
    configure_logging("DEBUG", stream)
    result = await summary.selected(cube_source["caller"], request)
    with cube_source["owner"].connect() as connection:
        rows = [
            dict(row)
            for row in connection.execute(
                select(call_log).where(call_log.c.user_id == cube_source["caller"].user_id)
            ).mappings()
        ]
    secrets = [sentence.text for sentence in result.sentences]
    safe = all(value not in stream.getvalue() + str(rows) for value in secrets)
    assert safe
    roots = [row for row in rows if row["name"] == "summarize.selected"]
    assert len(roots) == 1
    assert any(row["kind"] == "step" and row["parent_call_id"] is not None for row in rows)


def test_summarize_foreign_admin_and_disabled(cube_source, monkeypatch):
    from uuid import uuid4

    other = "msn-foreign-" + uuid4().hex[:20]
    with cube_source["owner"].begin() as connection:
        connection.execute(
            m.missions.insert().values(
                key=other,
                title="Simulated foreign mission",
                current_stage="st3-design",
                simulated=True,
            )
        )
    path = f"missions/{other}/engineering/st3-design/source.md"
    cube_source["provider"].write_content(path, b"# Simulated source\n\nA simulated foreign fact.")
    index_provider(cube_source["settings"])
    item = next(row for row in cube_source["provider"].items().values() if row["path"] == path)
    version = item["versions"][-1]
    request = {"item": item["item"], "version": version["version"], "sha256": version["sha256"]}
    called = []
    original = FakeChatClient.reply

    def spy(self, messages):
        called.append(True)
        return original(self, messages)

    monkeypatch.setattr(FakeChatClient, "reply", spy)
    try:
        with pytest.raises(NotFound):
            asyncio.run(summary.selected(cube_source["caller"], request))
        assert not called
        with TestClient(cube_source["app"], base_url="https://testserver") as http:
            assert cube_source["sign_in"](http).status_code == 200
            headers = {"Origin": "https://testserver"}
            assert http.post("/api/summarize", json=request, headers=headers).status_code == 404
            assert not called
            with cube_source["owner"].begin() as connection:
                connection.execute(
                    m.metadata.tables["user_roles"]
                    .insert()
                    .values(user_id=cube_source["caller"].user_id, role="admin")
                )
            assert http.post("/api/summarize", json=request, headers=headers).status_code == 200
            with cube_source["owner"].begin() as connection:
                users = m.metadata.tables["users"]
                connection.execute(
                    users.update()
                    .where(users.c.id == cube_source["caller"].user_id)
                    .values(disabled_at=datetime.now(UTC))
                )
            assert http.post("/api/summarize", json=request, headers=headers).status_code == 401
            assert len(called) == 1
    finally:
        with cube_source["owner"].begin() as connection:
            versions = list(
                connection.execute(
                    select(m.doc_versions.c.id).where(m.doc_versions.c.item == item["item"])
                ).scalars()
            )
            chunks = list(
                connection.execute(
                    select(m.chunks.c.id).where(m.chunks.c.doc_version_id.in_(versions))
                ).scalars()
            )
            connection.execute(
                m.vectors.delete().where(m.vectors.c.owner_id.in_([*versions, *chunks]))
            )
            connection.execute(m.doc_versions.delete().where(m.doc_versions.c.item == item["item"]))
            connection.execute(
                m.provider_items.delete().where(m.provider_items.c.item == item["item"])
            )
            facets = list(
                connection.execute(
                    select(m.facet_texts.c.id).where(m.facet_texts.c.mission_key == other)
                ).scalars()
            )
            connection.execute(
                m.vectors.delete().where(
                    m.vectors.c.owner_kind == "facet_text", m.vectors.c.owner_id.in_(facets)
                )
            )
            connection.execute(m.facet_texts.delete().where(m.facet_texts.c.mission_key == other))
            connection.execute(m.missions.delete().where(m.missions.c.key == other))


@run_async
async def test_model_rating_does_not_change_confidence(cube_source, monkeypatch):
    request = selected_source(cube_source)
    original = FakeChatClient.reply

    def rated(self, messages):
        result = original(self, messages)
        data = json.loads(result.text)
        data["selfRating"] = 1.0
        return ChatResponse(
            messages=[Message(role="assistant", contents=[json.dumps(data)])], model="fake"
        )

    monkeypatch.setattr(FakeChatClient, "reply", rated)
    result = await summary.selected(cube_source["caller"], request)
    assert result.confidence.score == 0 and result.confidence.level == "low"


@run_async
async def test_conflict_is_detected_without_model_enumeration(cube_source, monkeypatch):
    original = FakeChatClient.reply

    def omitted(self, messages):
        result = original(self, messages)
        data = json.loads(result.text)
        data["filings"] = [candidate for candidate in data["filings"] if candidate["value"] != "C"]
        return ChatResponse(
            messages=[Message(role="assistant", contents=[json.dumps(data)])], model="fake"
        )

    monkeypatch.setattr(FakeChatClient, "reply", omitted)
    result = await summary.upload(cube_source["caller"], upload_file("docx", conflict=True))
    assert result.stage.value is None and result.stage.confidence == 0
    assert len(result.stage.evidence) >= 2


def test_http_upload_and_failure_logging_chain(cube_source, monkeypatch, capsys, caplog):
    from app.infrastructure.operations.storage import call_log

    request = selected_source(cube_source)
    original = upload_file("pdf")
    output = io.StringIO()
    configure_logging("DEBUG", output)
    marker = hashlib.sha256(b"simulated failure marker").hexdigest()
    with TestClient(cube_source["app"], base_url="https://testserver") as http:
        assert cube_source["sign_in"](http).status_code == 200
        with cube_source["owner"].connect() as connection:
            before = connection.execute(select(func.max(call_log.c.id))).scalar_one() or 0
        headers = {"Origin": "https://testserver"}
        selected = http.post("/api/summarize", json=request, headers=headers)
        assert selected.status_code == 200
        uploaded = http.post(
            "/api/summarize/upload",
            content=original.content,
            headers=headers
            | {
                "Content-Type": original.media_type,
                "X-File-Name": original.file_name,
                "X-Simulated": "true",
            },
        )
        assert uploaded.status_code == 200

        def failure(self, messages):
            raise RuntimeError(marker)

        monkeypatch.setattr(FakeChatClient, "reply", failure)
        rejected = http.post("/api/summarize", json=request, headers=headers)
        assert rejected.status_code == 500 and rejected.json()["code"] == "internal_error"
        with cube_source["owner"].connect() as connection:
            rows = [
                dict(row)
                for row in connection.execute(
                    select(call_log).where(call_log.c.id > before).order_by(call_log.c.id)
                ).mappings()
            ]
        roots = [row for row in rows if row["name"] in {"summarize.selected", "summarize.upload"}]
        assert len(roots) == 3 and all(
            row["origin"] == "http"
            and row["auth_session_id"] is not None
            and row["user_id"] == cube_source["caller"].user_id
            for row in roots
        )
        for response, root in zip([selected, uploaded, rejected], roots, strict=True):
            assert str(root["request_id"]) == response.headers["x-request-id"]
            linked = [row for row in rows if row["request_id"] == root["request_id"]]
            ids = {row["call_id"] for row in linked}
            assert all(
                row["user_id"] == root["user_id"] and row["parent_call_id"] in ids
                for row in linked
                if row is not root and row["call_id"] != root["call_id"]
            )
    captured = capsys.readouterr()
    leaked = (
        output.getvalue()
        + captured.out
        + captured.err
        + str(rows)
        + " ".join(record.getMessage() for record in caplog.records)
    )
    secrets = [
        marker,
        *[sentence["text"] for sentence in selected.json()["sentences"]],
        *[sentence["text"] for sentence in uploaded.json()["sentences"]],
    ]
    assert all(secret not in leaked for secret in secrets)
