"""Synthetic files verify source order, locators, inactive cells and indexing."""

import re
import signal
import tracemalloc
import zlib
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile

import pytest

from app.library.convert import convert
from app.library.tests.files import (
    csv_file,
    pdf_file,
    slides_file,
    word_file,
    workbook_file,
)


def read(path: Path):
    return convert(path.read_bytes(), path.name)


def test_docx_headings_and_tables(tmp_path) -> None:
    blocks = read(word_file(tmp_path / "simulated.docx"))
    assert [block.kind for block in blocks] == ["text", "table", "text"]
    assert [block.heading_path for block in blocks] == [
        "1 Simulated purpose",
        "1 Simulated purpose",
        "2 Simulated design",
    ]
    assert all(
        block.locator.startswith("heading=") and ";para=" in block.locator for block in blocks
    )
    assert "before the table" in blocks[0].text and "after the table" in blocks[2].text


@pytest.mark.slow
def test_pdf_pages(tmp_path) -> None:
    blocks = read(pdf_file(tmp_path / "simulated.pdf"))
    assert [block.locator for block in blocks] == [
        "page=1;block=1",
        "page=2;block=1",
        "page=3;block=1",
    ]


def test_pdf_without_text_layer(tmp_path) -> None:
    with pytest.raises(ValueError, match="^no text layer$"):
        read(pdf_file(tmp_path / "simulated.pdf", text=False))


def test_pptx_slides_with_notes(tmp_path) -> None:
    blocks = read(slides_file(tmp_path / "simulated.pptx"))
    assert [block.locator for block in blocks] == ["slide=1", "slide=2"]
    assert [block.heading_path for block in blocks] == ["Simulated slide 1", "Simulated slide 2"]
    assert all(f"speaker note {number}" in block.text for number, block in enumerate(blocks, 1))


def test_xlsx_regions_values_only(tmp_path) -> None:
    blocks = read(workbook_file(tmp_path / "simulated.xlsx"))
    assert len(blocks) == 2 and all(block.kind == "table" for block in blocks)
    assert blocks[0].locator == "sheet=Power;range=A1:C3;rows=1-2"
    assert blocks[1].locator == "sheet=Power;range=A5:C6;rows=1-1"
    assert "=1/0" not in blocks[1].text and "| Formula |  | W |" in blocks[1].text


@pytest.mark.parametrize("format", ["docx", "xlsx", "csv"])
def test_long_table_split(tmp_path, format) -> None:
    path = tmp_path / ("simulated." + format)
    if format == "docx":
        word_file(path, rows=75)
    elif format == "xlsx":
        workbook_file(path, rows=75, second_region=False)
    else:
        csv_file(path, rows=75)
    blocks = [block for block in read(path) if block.kind == "table"]
    assert len(blocks) == 3
    header = blocks[0].text.splitlines()[0]
    assert [len(block.text.splitlines()) - 2 for block in blocks] == [30, 30, 15]
    assert all(block.text.splitlines()[0] == header for block in blocks)
    assert [block.locator.rsplit(";rows=", 1)[1] for block in blocks] == ["1-30", "31-60", "61-75"]


def test_csv(tmp_path) -> None:
    blocks = read(csv_file(tmp_path / "simulated.csv"))
    assert len(blocks) == 1 and blocks[0].locator == "sheet=simulated.csv;range=A1:C3;rows=1-2"


@pytest.mark.parametrize(
    "filename", ["simulated.xlsx", "simulated.docx", "simulated.pptx", "simulated.pdf"]
)
def test_mismatched_extension(filename) -> None:
    with pytest.raises(ValueError, match="^unsupported format$"):
        convert(b"Simulated plain text.", filename)


def test_office_package_type_must_match(tmp_path) -> None:
    with pytest.raises(ValueError, match="^unsupported format$"):
        convert(word_file(tmp_path / "simulated.docx").read_bytes(), "simulated.xlsx")


def test_paragraph_split_and_deterministic_handles(tmp_path) -> None:
    from docx import Document

    path = tmp_path / "simulated.docx"
    document = Document()
    document.add_heading("1 Synthetic heading", level=1)
    document.add_paragraph(" ".join(f"synthetic{n}" for n in range(650)))
    document.save(str(path))
    first = read(path)
    second = read(path)
    assert first == second and [len(block.text.split()) for block in first] == [300, 300, 50]
    assert all(block.locator == "heading=1;para=2" for block in first)


def test_multiline_and_pipe_cells_preserve_row_boundaries(tmp_path) -> None:
    from docx import Document

    path = word_file(tmp_path / "simulated.docx")
    document = Document(str(path))
    document.tables[0].cell(1, 0).text = "Mass|probe\npart"
    document.save(str(path))
    table = next(block for block in read(path) if block.kind == "table")
    assert len(table.text.splitlines()) == 4
    assert "Mass\\|probe part" in table.text


def test_invalid_pdf_diagnostics_never_log_source_bytes(tmp_path) -> None:
    import io
    import logging

    from app.infrastructure.logging import configure_logging

    stream = io.StringIO()
    configure_logging("DEBUG", stream)
    parser = logging.getLogger("pypdf")
    original_level = parser.level
    parser.setLevel(logging.DEBUG)
    marker = b"source_object_marker"
    payload = (
        pdf_file(tmp_path / "simulated.pdf").read_bytes().replace(b"1 0 obj", marker + b" obj", 1)
    )
    try:
        with pytest.raises(ValueError, match="^unsupported format$"):
            convert(payload, "simulated.pdf")
        leaked = marker.decode() in stream.getvalue()
        assert not leaked
    finally:
        parser.setLevel(original_level)


def test_pdf_table_repeats_headers_and_limits_rows(tmp_path) -> None:
    from fpdf import FPDF

    document = FPDF()
    document.add_page()
    document.set_font("Helvetica", size=7)
    for row in range(76):
        values = ["Parameter", "Value", "Unit"] if row == 0 else ["Mass", str(row), "kg"]
        for column, value in enumerate(values):
            document.set_xy(20 + column * 50, 20 + row * 2.5)
            document.cell(45, 2.5, value)
    path = tmp_path / "simulated.pdf"
    document.output(path)
    tables = [block for block in read(path) if block.kind == "table"]
    assert len(tables) == 3
    assert [len(block.text.splitlines()) - 2 for block in tables] == [30, 30, 15]
    assert all(block.locator.startswith("page=1;block=") for block in tables)
    assert all(block.text.splitlines()[0] == tables[0].text.splitlines()[0] for block in tables)


def test_slide_tables_follow_header_window_rule(tmp_path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(76, 3, Inches(1), Inches(1), Inches(6), Inches(4)).table
    for row in range(76):
        for column, value in enumerate(
            ["Parameter", "Value", "Unit"] if row == 0 else ["Mass", str(row), "kg"]
        ):
            table.cell(row, column).text = value
    path = tmp_path / "simulated.pptx"
    presentation.save(str(path))
    tables = [block for block in read(path) if block.kind == "table"]
    assert [len(block.text.splitlines()) - 2 for block in tables] == [30, 30, 15]
    assert all(block.locator.startswith("slide=1;table=1;rows=") for block in tables)


def test_html_cells_keep_literal_boundaries_and_one_row():
    source = (
        b"<table><tr><th>Mass (kg)</th><th>Label|literal</th></tr>"
        b"<tr><td>1|2</td><td>A<br>B<script>inactive</script></td></tr></table>"
    )
    tables = [block for block in convert(source, "synthetic.html") if block.kind == "table"]
    expected = [r"| Mass (kg) | Label\|literal |", "|---|---|", r"| 1\|2 | A B |"]
    preserved = len(tables) == 1 and tables[0].text.splitlines() == expected
    assert preserved
    assert tables[0].locator.endswith(";rows=1-1")


def test_html_optional_end_tags_keep_rows():
    source = b"<table><tr><th>Mass (kg)<th>Label<tr><td>1<td>A</table>"
    tables = [block for block in convert(source, "synthetic.html") if block.kind == "table"]
    expected = ["| Mass (kg) | Label |", "|---|---|", "| 1 | A |"]
    preserved = len(tables) == 1 and tables[0].text.splitlines() == expected
    assert preserved


def repack(path: Path, name: str, change) -> bytes:
    """The package at `path` with entry `name` changed."""
    out = BytesIO()
    with ZipFile(path) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info)
            target.writestr(info, change(data) if info.filename == name else data)
    return out.getvalue()


def corrupted(path: Path, name: str) -> bytes:
    """The package at `path` with entry `name`'s deflate data starting 0xFF: an invalid block."""
    content = bytearray(path.read_bytes())
    with ZipFile(path) as package:
        header = package.getinfo(name).header_offset
    # A 30-byte local header ends with the name's and the extra field's lengths, then the data.
    name_length = int.from_bytes(content[header + 26 : header + 28], "little")
    extra_length = int.from_bytes(content[header + 28 : header + 30], "little")
    content[header + 30 + name_length + extra_length] = 0xFF
    return bytes(content)


def hidden_tail(path: Path, name: str, tail: int) -> bytes:
    """The package with `tail` zero bytes deflated after entry `name`'s data, while both of its
    headers still give that data's own size and checksum."""
    with ZipFile(path) as source:
        entries = {info.filename: source.read(info) for info in source.infolist()}
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as target:
        for entry, data in entries.items():
            target.writestr(entry, data + bytes(tail) if entry == name else data)
    content = bytearray(out.getvalue())
    data = entries[name]
    honest = zlib.crc32(data).to_bytes(4, "little"), len(data).to_bytes(4, "little")
    with ZipFile(BytesIO(bytes(content))) as package:
        local = package.getinfo(name).header_offset  # its crc at +14, its size at +22
    content[local + 14 : local + 18], content[local + 22 : local + 26] = honest
    central = content.find(b"PK\x01\x02")  # a central record: crc at +16, size at +24
    while content[central + 46 : central + 46 + len(name)] != name.encode():
        central = content.find(b"PK\x01\x02", central + 4)
    content[central + 16 : central + 20], content[central + 24 : central + 28] = honest
    return bytes(content)


@contextmanager
def within(seconds: float):
    """Fail the block with TimeoutError once it runs past `seconds`, rather than hang."""

    def expire(signum, frame):
        raise TimeoutError(f"took more than {seconds} s")

    previous = signal.signal(signal.SIGALRM, expire)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def test_a_wrong_sheet_dimension_converts_as_before(tmp_path) -> None:
    path = workbook_file(tmp_path / "simulated.xlsx")
    wrong = repack(
        path,
        "xl/worksheets/sheet1.xml",
        lambda xml: re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1:XFD200000"/>', xml),
    )
    assert wrong != path.read_bytes() and convert(wrong, path.name) == read(path)


def test_a_sheet_spanning_a1_to_xfd200000_a_wide_csv_and_a_large_package_are_refused_quickly(
    tmp_path,
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"], sheet["XFD200000"] = "Parameter", "W"
    spread = tmp_path / "spread.xlsx"
    workbook.save(spread)
    # 1,001 columns by 1,999 rows pads to 2,000,999 cells; one row fewer would convert.
    wide = tmp_path / "wide.csv"
    wide.write_text("Parameter" + "," * 1000 + "\n" + "Mass\n" * 1998)
    large = word_file(tmp_path / "large.docx")
    with ZipFile(large, "a", ZIP_DEFLATED) as package:
        package.writestr("word/media/padding.bin", bytes(33 * 1024 * 1024))
    for path in (spread, wide, large):
        with within(5), pytest.raises(ValueError, match="^unsupported format$"):
            read(path)
    narrower = tmp_path / "narrower.csv"
    narrower.write_text("Parameter" + "," * 1000 + "\n" + "Mass\n" * 1997)
    assert read(narrower)


def test_data_hidden_after_an_entrys_declared_size_is_never_inflated(tmp_path) -> None:
    path = word_file(tmp_path / "simulated.docx")
    expected = read(path)
    lying = hidden_tail(path, "word/document.xml", 64 * 1024 * 1024)
    tracemalloc.start()
    try:
        assert convert(lying, path.name) == expected
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 16 * 1024 * 1024


def test_one_malformed_file_per_format_is_an_unsupported_format(tmp_path) -> None:
    """A cell span without its value, a slide without its shape tree, a cut-off sheet, a PDF
    whose root is no catalog and an unclosed CSV quote, plus a package in bzip2, a damaged deflate
    stream and a PDF stream in an unknown filter: each parser's error, never a crash."""
    slide = b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
    # The root is a page tree; /Size past pypdf's 10,000-object root search makes it give up
    # with LimitReachedError, a PyPdfError that is no PdfReadError.
    pdf = b"%PDF-1.4\n1 0 obj\n<</Type/Pages/Kids[]/Count 0>>\nendobj\n"
    pdf += (
        b"xref\n0 2\n0000000000 65535 f \n0000000009 00000 n \n"
        b"trailer\n<</Size 1000000/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % len(pdf)
    )
    malformed = {
        "simulated.docx": repack(
            word_file(tmp_path / "simulated.docx"),
            "word/document.xml",
            lambda xml: xml.replace(b"<w:tc>", b"<w:tc><w:tcPr><w:gridSpan/></w:tcPr>", 1),
        ),
        "simulated.pptx": repack(
            slides_file(tmp_path / "simulated.pptx"), "ppt/slides/slide1.xml", lambda _: slide
        ),
        "simulated.xlsx": repack(
            workbook_file(tmp_path / "simulated.xlsx"),
            "xl/worksheets/sheet1.xml",
            lambda xml: xml[: len(xml) // 2],
        ),
        "simulated.pdf": pdf,
        "simulated.csv": b'Parameter,"Value\n',
    }
    bzip2 = BytesIO()
    with (
        ZipFile(word_file(tmp_path / "bzip2.docx")) as source,
        ZipFile(bzip2, "w", ZIP_BZIP2) as target,
    ):
        for info in source.infolist():
            target.writestr(info.filename, source.read(info))
    damaged = {
        "bzip2.docx": bzip2.getvalue(),
        "damaged.docx": corrupted(word_file(tmp_path / "damaged.docx"), "word/document.xml"),
        "filtered.pdf": pdf_file(tmp_path / "filtered.pdf", pages=1)
        .read_bytes()
        .replace(b"/FlateDecode", b"/UnknownCode"),
    }
    for name, content in (malformed | damaged).items():
        with pytest.raises(ValueError, match="^unsupported format$"):
            convert(content, name)
