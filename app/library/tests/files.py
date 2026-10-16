"""Build small, explicitly synthetic Office fixtures in the supplied test folder."""

import csv
from datetime import UTC, datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

STAMP = datetime(2026, 1, 1, tzinfo=UTC)


def stable_package(path: Path) -> None:
    """Keep fixture package bytes independent of save time and file timestamps."""
    with ZipFile(path) as package:
        entries = {name: package.read(name) for name in package.namelist()}
    if "docProps/core.xml" in entries:
        core = ElementTree.fromstring(entries["docProps/core.xml"])
        for node in core:
            local = node.tag.rsplit("}", 1)[-1]
            if local in {"created", "modified"}:
                node.text = "2026-01-01T00:00:00Z"
            if local in {"creator", "lastModifiedBy"}:
                node.text = ""
        entries["docProps/core.xml"] = ElementTree.tostring(
            core, encoding="utf-8", xml_declaration=True
        )
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as package:
        for name, data in sorted(entries.items()):
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            package.writestr(info, data)


def word_file(path: Path, rows: int = 2) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("1 Simulated purpose", level=1)
    document.add_paragraph("Synthetic purpose text before the table.")
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Parameter", "Value", "Unit"], strict=True):
        cell.text = value
    for number in range(rows):
        for cell, value in zip(table.add_row().cells, ["Mass", str(number + 1), "kg"], strict=True):
            cell.text = value
    document.add_heading("2 Simulated design", level=1)
    document.add_paragraph("Synthetic design text after the table.")
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.save(str(path))
    stable_package(path)
    return path


def pdf_file(path: Path, pages: int = 3, text: bool = True) -> Path:
    from fpdf import FPDF

    document = FPDF()
    document.set_creation_date(STAMP)
    document.set_author("")
    document.set_creator("")
    for page in range(1, pages + 1):
        document.add_page()
        if text:
            document.set_font("Helvetica", size=12)
            document.multi_cell(0, 10, f"Simulated page {page}. Synthetic radiator margin.")
        else:
            document.rect(20, 20, 40, 40)
    document.output(path)
    return path


def slides_file(path: Path) -> Path:
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape

    presentation = Presentation()
    for number in [1, 2]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        title = slide.shapes.title
        assert title is not None
        title.text = f"Simulated slide {number}"
        body = slide.placeholders[1]
        assert isinstance(body, Shape)
        body.text = f"Synthetic slide body {number}."
        notes = slide.notes_slide.notes_text_frame
        assert notes is not None
        notes.text = f"Synthetic speaker note {number}."
    presentation.core_properties.author = ""
    presentation.core_properties.last_modified_by = ""
    presentation.save(str(path))
    stable_package(path)
    return path


def workbook_file(path: Path, rows: int = 2, second_region: bool = True) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Power"
    sheet.append(["Parameter", "Value", "Unit"])
    for number in range(rows):
        sheet.append(["Power", number + 1, "W"])
    if second_region:
        sheet.append([None, None, None])
        sheet.append(["Parameter", "Value", "Unit"])
        sheet.append(["Formula", "=1/0", "W"])
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.save(path)
    stable_package(path)
    return path


def csv_file(path: Path, rows: int = 2) -> Path:
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Parameter", "Value", "Unit"])
        for number in range(rows):
            writer.writerow(["Power", number + 1, "W"])
    return path


def office_files(root: Path) -> list[Path]:
    root.mkdir(parents=True, exist_ok=True)
    return [
        word_file(root / "simulated.docx"),
        pdf_file(root / "simulated.pdf"),
        slides_file(root / "simulated.pptx"),
        workbook_file(root / "simulated.xlsx"),
        csv_file(root / "simulated.csv"),
    ]
