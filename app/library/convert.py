"""Text converters preserve block order, headings and stable source locators."""

import csv
import logging
import re
import zlib
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import PurePath
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

# Untrusted files stay bounded: an Office package unpacks to at most MAX_UNPACKED bytes, and a
# workbook or CSV holds at most MAX_CELLS cells once every row is as wide as its sheet's widest.
MAX_UNPACKED = 32 * 1024 * 1024
MAX_CELLS = 2_000_000


@dataclass(frozen=True)
class Block:
    block: int
    kind: str
    heading_path: str
    locator: str
    text: str


class VisibleHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts = []
        self.hidden = 0
        self.first_row = False
        self.row_cells = 0
        self.table_depth = 0
        self.cell: list[str] | None = None
        self.row_open = False

    def flush_cell(self):
        if self.cell is not None:
            value = " ".join("".join(self.cell).split())
            value = value.replace("\\", "\\\\").replace("|", "\\|")
            self.parts.append(("| " if self.row_cells == 1 else " | ") + value)
            self.cell = None

    def flush_row(self):
        if self.row_open:
            self.flush_cell()
            if self.row_cells:
                self.parts.append(" |\n")
                if self.first_row:
                    self.parts.append("|" + "---|" * self.row_cells + "\n")
                    self.first_row = False
            self.row_open = False

    def start_row(self):
        self.flush_row()
        if self.parts and not self.parts[-1].endswith("\n"):
            self.parts.append("\n")
        self.row_cells = 0
        self.row_open = True

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        elif not self.hidden:
            if tag == "table":
                self.table_depth += 1
                if self.table_depth == 1:
                    self.parts.append("\n")
                    self.first_row = True
                elif self.cell is not None:
                    self.cell.append(" ")
            elif self.table_depth == 1 and tag in {"td", "th"}:
                if not self.row_open:
                    self.start_row()
                self.flush_cell()
                self.row_cells += 1
                self.cell = []
            elif self.table_depth == 1 and tag == "tr":
                self.start_row()
            elif self.cell is not None:
                if tag in {"p", "div", "br", "tr", "td", "th"} or re.fullmatch(r"h[1-6]", tag):
                    self.cell.append(" ")
            elif re.fullmatch(r"h[1-6]", tag):
                self.parts.append("\n" + "#" * int(tag[1]) + " ")
            elif tag in {"p", "div", "br"}:
                self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden:
            if self.table_depth == 1 and tag in {"td", "th"}:
                self.flush_cell()
            elif self.table_depth == 1 and tag == "tr":
                self.flush_row()
            elif tag == "table":
                if self.table_depth == 1:
                    self.flush_row()
                self.table_depth = max(0, self.table_depth - 1)
                if self.cell is not None:
                    self.cell.append(" ")
            elif self.cell is not None:
                if tag in {"p", "div", "tr", "td", "th"}:
                    self.cell.append(" ")
            elif tag in {"p", "div"} or re.fullmatch(r"h[1-6]", tag):
                self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            if self.cell is not None:
                self.cell.append(data)
            else:
                self.parts.append(data)


def text_parts(text: str, size: int = 300):
    words = list(re.finditer(r"\S+", text))
    for start in range(0, len(words), size):
        end = min(start + size, len(words))
        yield text[words[start].start() : words[end - 1].end()]


def text_document(content: bytes, filename: str) -> list[Block]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix not in {"md", "txt", "html", "htm"}:
        raise ValueError("unsupported format")
    text = content.decode("utf-8")
    if suffix in {"html", "htm"}:
        parser = VisibleHTML()
        parser.feed(text)
        text = "".join(parser.parts)
    blocks = []
    headings = []
    buffer = []
    heading_id = "0"

    def add(value, kind="text", rows=None):
        number = len(blocks) + 1
        locator = f"heading={heading_id};block={number}" + (f";rows={rows}" if rows else "")
        blocks.append(Block(number, kind, " > ".join(t for _, t in headings), locator, value))

    def flush():
        nonlocal buffer
        for part in text_parts("\n".join(buffer)):
            add(part)
        buffer = []

    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        heading = re.match(r"^(#{1,6})\s+(.+)$", line) if suffix != "txt" else None
        if heading:
            flush()
            level = len(heading[1])
            headings[:] = [(depth, title) for depth, title in headings if depth < level]
            headings.append((level, heading[2]))
            heading_id = heading[2].split()[0]
        elif (
            "|" in line and index + 1 < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[index + 1])
        ):
            flush()
            header = line
            separator = lines[index + 1]
            rows = []
            index += 2
            while index < len(lines) and "|" in lines[index]:
                rows.append(lines[index])
                index += 1
            for start in range(0, len(rows), 30):
                end = min(start + 30, len(rows))
                add("\n".join([header, separator, *rows[start:end]]), "table", f"{start + 1}-{end}")
            index -= 1
        else:
            buffer.append(line)
        index += 1
    flush()
    return blocks


def cell_text(value) -> str:
    return "" if value is None else " ".join(str(value).splitlines()).replace("\t", " ")


def add_table(blocks: list[Block], rows: list[list[str]], heading: str, locator) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    normalized = [
        [cell_text(value).replace("\\", "\\\\").replace("|", "\\|") for value in row]
        + [""] * (width - len(row))
        for row in rows
    ]

    def line(row):
        return "| " + " | ".join(row) + " |"

    header = line(normalized[0])
    separator = "|" + "---|" * width
    data = normalized[1:]
    for start in range(0, max(1, len(data)), 30):
        end = min(start + 30, len(data))
        source = locator(start, end)
        if data:
            source += f";rows={start + 1}-{end}"
        blocks.append(
            Block(
                len(blocks) + 1,
                "table",
                heading,
                source,
                "\n".join([header, separator, *map(line, data[start:end])]),
            )
        )


def word_document(content: bytes) -> list[Block]:
    from docx import Document
    from docx.table import Table

    document = Document(BytesIO(content))
    blocks = []
    headings = []
    heading_id = "0"
    paragraphs = []
    paragraph_number = 0
    table_number = 0

    def flush():
        combined = "\n".join(value for _, value in paragraphs)
        offset = 0
        for part in text_parts(combined):
            start = combined.find(part, offset)
            cursor = 0
            source = paragraphs[0][0]
            for number, value in paragraphs:
                if cursor + len(value) > start:
                    source = number
                    break
                cursor += len(value) + 1
            blocks.append(
                Block(
                    len(blocks) + 1,
                    "text",
                    " > ".join(t for _, t in headings),
                    f"heading={heading_id};para={source}",
                    part,
                )
            )
            offset = start + len(part)
        paragraphs.clear()

    for entry in document.iter_inner_content():
        if isinstance(entry, Table):
            flush()
            table_number += 1
            base = f"heading={heading_id};para={paragraph_number};table={table_number}"
            add_table(
                blocks,
                [[cell.text for cell in row.cells] for row in entry.rows],
                " > ".join(t for _, t in headings),
                lambda start, end, base=base: base,
            )
            continue
        paragraph_number += 1
        style = entry.style
        match = re.fullmatch(r"heading\s*([1-6])", (style.name or "").lower()) if style else None
        if match:
            flush()
            level = int(match[1])
            headings[:] = [(depth, title) for depth, title in headings if depth < level]
            headings.append((level, entry.text))
            number = re.match(r"^\d+(?:\.\d+)*", entry.text)
            heading_id = number[0] if number else entry.text
        elif entry.text.strip():
            paragraphs.append((paragraph_number, entry.text))
    flush()
    return blocks


def pdf_document(content: bytes) -> list[Block]:
    from pypdf import PdfReader

    # Parser diagnostics can include supplied object bytes; expose only fixed errors.
    logging.getLogger("pypdf").setLevel(logging.CRITICAL + 1)
    reader = PdfReader(BytesIO(content), strict=True)
    pages = [
        page.extract_text(extraction_mode="layout", layout_mode_space_vertically=False) or ""
        for page in reader.pages
    ]
    if not any(page.strip() for page in pages):
        raise ValueError("no text layer")
    blocks = []
    for page_number, text in enumerate(pages, 1):
        heading = f"Page {page_number}"
        lines = text.splitlines()
        buffer = []
        page_block = 0

        def flush(buffer=buffer, heading=heading, page_number=page_number):
            nonlocal page_block
            for part in text_parts("\n".join(buffer)):
                page_block += 1
                blocks.append(
                    Block(
                        len(blocks) + 1,
                        "text",
                        heading,
                        f"page={page_number};block={page_block}",
                        part,
                    )
                )
            buffer.clear()

        index = 0
        while index < len(lines):
            cells = re.split(r"\t+| {2,}", lines[index].strip())
            region = [cells]
            end = index + 1
            if len(cells) > 1:
                while end < len(lines):
                    next_cells = re.split(r"\t+| {2,}", lines[end].strip())
                    if len(next_cells) != len(cells):
                        break
                    region.append(next_cells)
                    end += 1
            numeric = any(re.search(r"\d", cell) for row in region[1:] for cell in row)
            if len(region) >= 2 and numeric:
                flush()

                def table_locator(start, end, page_number=page_number):
                    nonlocal page_block
                    page_block += 1
                    return f"page={page_number};block={page_block}"

                add_table(blocks, region, heading, table_locator)
                index = end
            else:
                buffer.append(lines[index])
                index += 1
        flush()
    return blocks


def slides_document(content: bytes) -> list[Block]:
    from pptx import Presentation
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.graphfrm import GraphicFrame

    presentation = Presentation(BytesIO(content))
    blocks = []
    for number, slide in enumerate(presentation.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title else f"Slide {number}"
        texts = [shape.text for shape in slide.shapes if isinstance(shape, Shape)]
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                texts.append(frame.text)
        # Slides keep their notes together; explicit tables follow in slide order.
        blocks.append(Block(len(blocks) + 1, "slide", title, f"slide={number}", "\n".join(texts)))
        for table_number, shape in enumerate(
            (s for s in slide.shapes if isinstance(s, GraphicFrame) and s.has_table), 1
        ):
            base = f"slide={number};table={table_number}"
            add_table(
                blocks,
                [[cell.text for cell in row.cells] for row in shape.table.rows],
                title,
                lambda start, end, base=base: base,
            )
    return blocks


def sheet_regions(blocks: list[Block], rows: list[list[str]], name: str) -> None:
    from openpyxl.utils import get_column_letter

    start = 0
    while start < len(rows):
        if not any(rows[start]):
            start += 1
            continue
        end = start + 1
        while end < len(rows) and any(rows[end]):
            end += 1
        region = rows[start:end]
        used = [column for row in region for column, value in enumerate(row) if value != ""]
        first, last = min(used), max(used)
        trimmed = [(row + [""] * (last + 1 - len(row)))[first : last + 1] for row in region]
        header_row = start + 1

        def locator(low, high, header_row=header_row, first=first, last=last, end=end):
            # The range identifies the source region; rows identifies the window.
            return (
                f"sheet={name};range={get_column_letter(first + 1)}{header_row}:"
                f"{get_column_letter(last + 1)}{end}"
            )

        add_table(blocks, trimmed, name, locator)
        start = end


def workbook_document(content: bytes) -> list[Block]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    blocks = []
    cells = 0
    try:
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()  # a wrong <dimension> would pad every row to its range
            rows, width = [], 0
            for row in sheet.iter_rows(values_only=True):
                rows.append([cell_text(value) for value in row])
                width = max(width, len(row))
                padded(rows, width, cells)  # refused while reading, before rows pile up
            cells = padded(rows, width, cells)
            sheet_regions(blocks, rows, sheet.title)
    finally:
        workbook.close()
    return blocks


def csv_document(content: bytes, filename: str) -> list[Block]:
    text = content.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows, width = [], 0
    for row in csv.reader(StringIO(text), dialect, strict=True):
        rows.append(row)
        width = max(width, len(row))
        padded(rows, width)  # refused while reading, before rows pile up
    blocks = []
    sheet_regions(blocks, rows, PurePath(filename).name)
    return blocks


def padded(rows: list[list[str]], width: int, counted: int = 0) -> int:
    """`counted` plus these rows' cells with every row, empty or not, as wide as the widest (at
    least one cell). More than MAX_CELLS is refused."""
    cells = counted + len(rows) * max(width, 1)
    if cells > MAX_CELLS:
        raise ValueError("unsupported format")
    return cells


def unpacked(content: bytes) -> bytes:
    """The Office package stored uncompressed, refused past MAX_UNPACKED bytes. Each entry is read
    a megabyte at a time and never past the size it declares, so data hidden after that is never
    inflated; compressions other than deflate, and encryption, are refused."""
    copy, total = BytesIO(), 0
    with ZipFile(BytesIO(content)) as package, ZipFile(copy, "w", ZIP_STORED) as stored:
        for info in package.infolist():
            if info.compress_type not in (ZIP_STORED, ZIP_DEFLATED) or info.flag_bits & 0x1:
                raise ValueError("unsupported format")
            with package.open(info) as source, stored.open(info.filename, "w") as target:
                while chunk := source.read(1024 * 1024):
                    total += len(chunk)
                    if total > MAX_UNPACKED:
                        raise ValueError("unsupported format")
                    target.write(chunk)
    return copy.getvalue()


def convert(content: bytes, filename: str) -> list[Block]:
    """Read only the declared format; no Office formula, link, macro or script runs."""
    from docx.exceptions import PythonDocxError
    from docx.opc.exceptions import OpcError
    from docx.oxml.exceptions import XmlchemyError
    from lxml.etree import XMLSyntaxError
    from pptx.exc import PythonPptxError
    from pypdf.errors import PyPdfError

    suffix = PurePath(filename).suffix.lower()
    office = {
        ".docx": ("word/document.xml", word_document),
        ".pptx": ("ppt/presentation.xml", slides_document),
        ".xlsx": ("xl/workbook.xml", workbook_document),
    }
    try:
        if suffix in office:
            entry, reader = office[suffix]
            with ZipFile(BytesIO(content)) as package:
                if (
                    "[Content_Types].xml" not in package.namelist()
                    or entry not in package.namelist()
                ):
                    raise ValueError("unsupported format")
            return reader(unpacked(content))
        if suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("unsupported format")
            return pdf_document(content)
        if suffix == ".csv":
            return csv_document(content, filename)
        if suffix in {".md", ".txt", ".html", ".htm"}:
            if content.startswith((b"PK\x03\x04", b"%PDF-")) or b"\x00" in content:
                raise ValueError("unsupported format")
            return text_document(content, filename)
    except ValueError as error:
        if str(error) == "no text layer":
            raise
        raise ValueError("unsupported format") from None
    except (
        BadZipFile,
        UnicodeDecodeError,
        csv.Error,
        KeyError,
        RecursionError,
        NotImplementedError,
        zlib.error,
        ParseError,
        XMLSyntaxError,
        PyPdfError,
        PythonPptxError,
        PythonDocxError,
        OpcError,
        XmlchemyError,
    ):
        raise ValueError("unsupported format") from None
    raise ValueError("unsupported format")
