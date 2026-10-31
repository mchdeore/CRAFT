"""Focused generated-data checks of SDK, owned paths and server bindings."""

import asyncio
from types import SimpleNamespace
from typing import Any

import pytest
from mcp.types import CallToolResult, ListToolsResult, TextContent


def test_sdk_forwards_existing_names_and_arguments_once():
    from craft import Craft

    class Session:
        def __init__(self):
            self.requests = []

        async def list_tools(self):
            return ListToolsResult(tools=[])

        async def call_tool(self, name, arguments=None, *, meta=None):
            self.requests.append((name, arguments, meta))
            return CallToolResult(
                content=[TextContent(type="text", text='{"hits":[]}')],
                structuredContent={"hits": []},
            )

    async def check():
        session = Session()
        sdk = Craft(session)
        result = await sdk.search("generated source", limit=2)
        assert result.structuredContent == {"hits": []}
        assert len(session.requests) == 1
        name, arguments, meta = session.requests[0]
        assert (name, arguments) == ("search", {"query": "generated source", "limit": 2})
        assert isinstance(meta["craft/toolCallId"], str)
        await sdk.document("itm_012345abcdef", version=1)
        await sdk.skills(limit=10)
        assert [(name, args) for name, args, _ in session.requests[1:]] == [
            ("get_document", {"item": "itm_012345abcdef", "version": 1}),
            ("list_skills", {"limit": 10}),
        ]

        class LostResponse(Session):
            async def call_tool(self, name, arguments=None, *, meta=None):
                self.requests.append((name, arguments, meta))
                raise ConnectionError("response lost")

        lost = LostResponse()
        with pytest.raises(ConnectionError):
            await Craft(lost)._call("make_chart", {"spec": {}, "format": "png"})
        assert len(lost.requests) == 1

    asyncio.run(check())


def test_owned_path_rejects_escape_and_symlink(tmp_path):
    from craft import owned_path

    root = tmp_path / "owner"
    root.mkdir()
    (root / "scripts").mkdir()
    assert owned_path("scripts/a.py", root=root) == root / "scripts/a.py"
    outside = tmp_path / "outside"
    outside.mkdir()
    (root / "link").symlink_to(outside, target_is_directory=True)
    (root / "dangling").symlink_to(outside / "missing")
    for path in ["../outside/a.py", "/etc/passwd", "link/a.py", "dangling", "a/../b", ""]:
        with pytest.raises(ValueError):
            owned_path(path, root=root)


def test_official_mcp_connection_forwards_bearer_outside_arguments(monkeypatch):
    import json

    import httpx
    from craft import client as sdk_module
    from craft.client import connect

    actual_client = httpx.AsyncClient
    requests = []

    def respond(request):
        body = json.loads(request.content)
        requests.append((request, body))
        method = body["method"]
        if "id" not in body:
            return httpx.Response(202)
        if method == "initialize":
            result = {
                "protocolVersion": "2025-06-18",
                "capabilities": {"tools": {}},
                "serverInfo": {"name": "generated-test", "version": "1"},
            }
        elif method == "tools/list":
            result = {
                "tools": [
                    {
                        "name": "search",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "query": {"type": "string"},
                                "limit": {"type": "integer"},
                            },
                        },
                    }
                ]
            }
        elif method == "tools/call":
            result = {"content": [], "structuredContent": {"hits": []}, "isError": False}
        else:
            raise AssertionError(f"Unexpected fake request: {method}")
        return httpx.Response(200, json={"jsonrpc": "2.0", "id": body["id"], "result": result})

    def configured_client(**kwargs):
        return actual_client(transport=httpx.MockTransport(respond), **kwargs)

    monkeypatch.setattr(sdk_module, "httpx", SimpleNamespace(AsyncClient=configured_client))

    async def check():
        async with connect("http://127.0.0.1:9999/mcp", run_token="synthetic-owned-run") as sdk:
            result = await sdk.search("generated evidence", limit=1)
            assert result.structuredContent == {"hits": []}

    asyncio.run(check())
    calls = [(request, body) for request, body in requests if body["method"] == "tools/call"]
    assert len(calls) == 1
    request, body = calls[0]
    assert request.headers["authorization"] == "Bearer synthetic-owned-run"
    assert "cookie" not in request.headers
    assert body["params"]["arguments"] == {"query": "generated evidence", "limit": 1}
    assert "synthetic-owned-run" not in json.dumps(body)


def test_word_edit_roundtrip_preserves_original(tmp_path):
    from shutil import copyfile

    from craft import owned_path
    from docx import Document

    workspace = tmp_path / "owner"
    workspace.mkdir()
    original = owned_path("imports/source.docx", root=workspace)
    output = owned_path("outputs/edited.docx", root=workspace)
    original.parent.mkdir(parents=True)
    generated = Document()
    generated.add_paragraph("Generated source paragraph.")
    generated.add_paragraph("Keep this paragraph.")
    table = generated.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("Metric", "Value"), ("Revenue", "125"), ("Cost", "50")], strict=True
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    generated.save(str(original))
    original_bytes = original.read_bytes()

    output.parent.mkdir(parents=True)
    copyfile(original, output)
    edited = Document(str(output))
    assert [paragraph.text for paragraph in edited.paragraphs] == [
        "Generated source paragraph.",
        "Keep this paragraph.",
    ]
    assert [[cell.text for cell in row.cells] for row in edited.tables[0].rows] == [
        ["Metric", "Value"],
        ["Revenue", "125"],
        ["Cost", "50"],
    ]
    edited.paragraphs[0].text = "Edited paragraph in personal copy."
    edited.tables[0].cell(1, 1).text = "150"
    edited.add_paragraph("Summary: revenue increased by 25.")
    edited.save(str(output))

    saved = Document(str(output))
    assert [paragraph.text for paragraph in saved.paragraphs] == [
        "Edited paragraph in personal copy.",
        "Keep this paragraph.",
        "Summary: revenue increased by 25.",
    ]
    assert [[cell.text for cell in row.cells] for row in saved.tables[0].rows] == [
        ["Metric", "Value"],
        ["Revenue", "150"],
        ["Cost", "50"],
    ]
    assert original.read_bytes() == original_bytes


def test_table_finds_header_below_title_rows_and_path_accepts_its_own_result(tmp_path, monkeypatch):
    import pandas as pd
    from craft import _root, files

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    rows = [
        ["Simulated mission 02", "SIMULATED / UNCLASSIFIED", None, None],
        ["Valid from", "2026-01-01", None, None],
        [None, None, None, None],
        ["Element", "Masse sèche (g)", "Volume utile (L)", "Design life"],
        ["Mission total", 104000, 215, 5],
        ["payload", 46800, 96.75, 5],
    ]
    pd.DataFrame(rows).to_excel(tmp_path / "budget.xlsx", header=False, index=False)
    frame = files.table("budget.xlsx")
    assert list(frame.columns) == ["Element", "Masse sèche (g)", "Volume utile (L)", "Design life"]
    assert frame["Masse sèche (g)"].tolist() == [104000, 46800]
    assert pd.api.types.is_numeric_dtype(frame["Volume utile (L)"])
    saved = files.path("outputs/result.csv")
    assert files.path(saved) == saved
    assert files.save(frame, saved) == "outputs/result.csv"
    with pytest.raises(ValueError):
        files.path(tmp_path / ".." / "outside.csv")


def test_vega_chart_renders_offline_and_saves_spec_and_rows(tmp_path, monkeypatch):
    import json

    import pandas as pd
    from craft import _root, analysis

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    frame = pd.DataFrame({"mission": ["a", "b"], "power": [120.5, 80.0]})
    shown = analysis.chart(frame, "mission", "power", kind="bar", engine="vega")
    assert isinstance(shown, analysis.VegaChart)
    assert shown._repr_svg_().lstrip().startswith("<svg")
    saved = analysis.chart(
        frame, "mission", "power", kind="bar", engine="vega", save_as="outputs/charts/power.svg"
    )
    assert saved == "outputs/charts/power.svg"
    spec = json.loads((tmp_path / "outputs/charts/power.vl.json").read_text())
    assert spec["mark"] == "bar" and spec["data"]["values"][0] == {"mission": "a", "power": 120.5}
    assert (tmp_path / "outputs/charts/power.csv").read_text() == "mission,power\na,120.5\nb,80.0\n"
    assert analysis.vega(spec, save_as="power.png") == "power.png"
    with pytest.raises(ValueError):
        analysis.vega({**spec, "data": {"url": "https://example.invalid/data.json"}})


def _word(path, *paragraphs, header=None):
    from docx import Document

    document = Document()
    for runs in paragraphs:
        paragraph = document.add_paragraph()
        for text, bold in runs:
            paragraph.add_run(text).bold = bold
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    document.save(path)


def test_replace_in_word_matches_across_runs_and_in_headers(tmp_path, monkeypatch):
    from craft import _root, docs

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    _word(
        tmp_path / "memo.docx",
        [("Mission ", False), ("Ze", True), ("ta", False)],
        header="Zeta memo",
    )
    assert docs.replace_in_word("memo.docx", "Zeta", "Kappa", save_as="memo-2.docx") == 2
    text = docs.read_word("memo-2.docx")
    assert "Mission Kappa" in text
    from docx import Document

    edited = Document(str(tmp_path / "memo-2.docx"))
    assert edited.paragraphs[0].runs[1].bold and edited.paragraphs[0].runs[1].text == "Kappa"
    assert edited.sections[0].header.paragraphs[0].text == "Kappa memo"
    assert "Zeta" in docs.read_word("memo.docx")


def test_fill_template_word_excel_and_powerpoint_keep_the_template(tmp_path, monkeypatch):
    from craft import _root, docs
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    from pptx import Presentation

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    _word(
        tmp_path / "memo.docx",
        [("To: {{ recipient }}", False)],
        [("{% for r in rows %}{{ r }};{% endfor %}", False)],
    )
    original = (tmp_path / "memo.docx").read_bytes()
    assert docs.template_fields("memo.docx") == ["recipient", "rows"]
    with pytest.raises(ValueError, match="rows"):
        docs.fill_template("memo.docx", {"recipient": "Finance"}, save_as="out/memo.docx")
    saved = docs.fill_template(
        "memo.docx", {"recipient": "Finance & Ops", "rows": ["a", "b"]}, save_as="out/memo.docx"
    )
    assert saved == "out/memo.docx" and (tmp_path / "memo.docx").read_bytes() == original
    text = docs.read_word(saved)
    assert "To: Finance & Ops" in text and "a;b;" in text and "{{" not in text

    book: Any = Workbook()
    book.active.title = "Cost"
    book.defined_names["total"] = DefinedName("total", attr_text="Cost!$B$2")
    book.save(tmp_path / "cost.xlsx")
    assert docs.template_fields("cost.xlsx") == ["total"]
    docs.fill_template("cost.xlsx", {"total": 42}, save_as="cost-filled.xlsx")
    assert load_workbook(tmp_path / "cost-filled.xlsx")["Cost"]["B2"].value == 42

    deck = Presentation()
    slide: Any = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Review of {{ mission }}"
    deck.save(str(tmp_path / "deck.pptx"))
    assert docs.template_fields("deck.pptx") == ["mission"]
    docs.fill_template("deck.pptx", {"mission": "msn-sim-02"}, save_as="deck-filled.pptx")
    assert docs.read_slides("deck-filled.pptx")[0]["title"] == "Review of msn-sim-02"
    with pytest.raises(ValueError):
        docs.fill_template("deck.pptx", {"mission": "x"}, save_as="deck.pptx")


def test_workbook_and_slides_get_native_charts_from_a_vega_lite_spec(tmp_path, monkeypatch):
    import pandas as pd
    from craft import _root, docs
    from openpyxl import load_workbook
    from pptx import Presentation

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    frame = pd.DataFrame({"mission": ["a", "b"], "power": [120.5, 80]})
    spec = {
        "title": "Peak power",
        "data": {"values": frame.to_dict(orient="records")},
        "mark": "bar",
        "encoding": {"x": {"field": "mission"}, "y": {"field": "power"}},
    }
    assert docs.workbook("book.xlsx", ("Power", frame), charts={"Power": spec}) == "book.xlsx"
    sheet: Any = load_workbook(tmp_path / "book.xlsx")["Power"]
    assert sheet["B3"].value == 80 and len(sheet._charts) == 1
    docs.chart_slide("deck.pptx", spec)
    docs.chart_slide("deck.pptx", {**spec, "mark": "line"}, title="Trend")
    slides = Presentation(str(tmp_path / "deck.pptx")).slides
    assert len(slides) == 2 and any(shape.has_chart for shape in slides[1].shapes)


@pytest.mark.slow
def test_pages_and_preview_render_with_libreoffice_in_the_workspace_image():
    import subprocess

    # Guarded like the kernel, so page rendering is checked where the agent runs it.
    script = """
import asyncio, os
import craft_guard
assert craft_guard.installed()
from docx import Document
import craft
import craft_preview
document = Document(); document.add_heading("Preview", 0); document.add_paragraph("Generated.")
document.save("/workspace/memo.docx")
pdf = asyncio.run(craft_preview.preview("memo.docx"))
assert pdf.startswith(b"%PDF") and asyncio.run(craft_preview.preview("memo.docx")) == pdf
pages = craft.preview("memo.docx")
assert len(pages) == 1 and pages[0]._repr_png_().startswith(b"\\x89PNG")
report = craft.check()
assert all(report[part].startswith("ok") for part in
           ("LibreOffice", "Poppler", "Fonts", "Python libraries")), report
print("ok")
"""
    # The image scripts/check full builds, offline with a read-only root.
    run = subprocess.run(
        ["docker", "run", "--rm", "--network", "none", "--read-only", "--tmpfs", "/tmp"]
        + ["--tmpfs", "/workspace:uid=1000", "--env", "CRAFT_GUARD=1"]
        + ["craft-user-workspace:local-pilot", "python", "-c", script],
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert run.returncode == 0 and run.stdout.strip().endswith("ok"), run.stderr[-2000:]


def test_a_failed_save_keeps_the_previous_file_and_leaves_nothing_behind(tmp_path, monkeypatch):
    from craft import _root, files

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    assert files.save({"a": 1}, "out/data.json") == "out/data.json"

    class Broken:
        def __str__(self):
            raise RuntimeError("serializer failed")

    with pytest.raises(RuntimeError):
        files.save(Broken(), "out/data.json")
    assert (tmp_path / "out/data.json").read_text() == '{\n  "a": 1\n}\n'
    assert [p.name for p in (tmp_path / "out").iterdir()] == ["data.json"]


def test_the_preview_refuses_an_oversized_source_without_converting(tmp_path, monkeypatch):
    import craft_preview
    from craft import _root

    monkeypatch.setattr(_root, "USER_ROOT", tmp_path)
    monkeypatch.setattr(craft_preview, "MAX_BYTES", 10)
    (tmp_path / "big.docx").write_bytes(b"x" * 11)
    with pytest.raises(craft_preview.TooLarge):
        asyncio.run(craft_preview.preview("big.docx"))
