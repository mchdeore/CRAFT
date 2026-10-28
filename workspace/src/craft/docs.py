"""Read, edit and build Word, PowerPoint, PDF and Excel files in the workspace.

Names are relative to ``USER_ROOT``. Source documents from the CRAFT library
are read-only; import a personal copy before editing it.
"""

from pathlib import Path
from typing import Any

from craft import _root, files

_EDITABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm", ".docx", ".pptx", ".ipynb", ".md", ".txt"}

# Reading


def read_word(name: str | Path) -> str:
    """A Word document as Markdown: headings, paragraphs, lists and tables in order."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(files.path(name)))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                lines.append("#" * int(style[-1]) + " " + text)
            elif style == "Title":
                lines.append("# " + text)
            elif "List" in style:
                lines.append("- " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            rows = [
                [cell.text.strip().replace("|", "\\|") for cell in row.cells]
                for row in Table(child, document).rows
            ]
            if rows:
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "---|" * len(rows[0]))
                lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n\n".join(lines) + "\n"


def word_tables(name: str | Path) -> list:
    """Every Word table as a DataFrame, using its first row as the header."""
    import pandas as pd
    from docx import Document

    frames = []
    for table in Document(str(files.path(name))).tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            frames.append(pd.DataFrame(rows[1:], columns=pd.Index(_unique(rows[0]))))
    return frames


def read_slides(name: str | Path) -> list[dict[str, Any]]:
    """Each slide's number, title, text and speaker notes."""
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(files.path(name))).slides, start=1):
        heading: Any = slide.shapes.title
        title = heading.text if heading is not None else ""
        shapes: list[Any] = list(slide.shapes)
        texts = [
            shape.text_frame.text
            for shape in shapes
            if shape.has_text_frame and shape != heading and shape.text_frame.text
        ]
        notes_frame: Any = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = notes_frame.text if notes_frame is not None else ""
        slides.append({"slide": number, "title": title, "text": "\n".join(texts), "notes": notes})
    return slides


def read_pdf(name: str | Path) -> list[str]:
    """The text of each page. Scanned pages without a text layer come back empty."""
    from pypdf import PdfReader

    return [page.extract_text() or "" for page in PdfReader(str(files.path(name))).pages]


def sheets(name: str | Path) -> dict:
    """Every sheet of a workbook as a DataFrame, by sheet name. Formulas are not run."""
    return files.read(name, sheet=None)


# Editing


def editable_copy(name: str | Path, *, save_as: str | Path | None = None) -> str:
    """Copy a document or table file to a new editable derivative name.

    The original file is never modified. If ``save_as`` is omitted, this creates
    ``-edit`` (or ``-edit-2``, ``-edit-3``...) beside the source file.
    """
    source = files.path(name)
    suffix = source.suffix.lower()
    if suffix not in _EDITABLE_SUFFIXES:
        allowed = ", ".join(sorted(_EDITABLE_SUFFIXES))
        raise ValueError(f"Editable copies support {allowed}")
    content = source.read_bytes()
    if save_as is None:
        base = f"{source.stem}-edit"
        index = 1
        while True:
            stem = base if index == 1 else f"{base}-{index}"
            target = source.with_name(f"{stem}{source.suffix}")
            # Exclusive creation also refuses dangling symlinks and racing writers.
            try:
                with target.open("xb") as output:
                    output.write(content)
                break
            except FileExistsError:
                index += 1
    else:
        target = files.path(save_as)
        if target.suffix.lower() != suffix:
            raise ValueError("Editable copies keep the same file extension")
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("xb") as output:
            output.write(content)
    return target.relative_to(_root.USER_ROOT).as_posix()


def open_word(name: str | Path):
    """A python-docx Document to edit; save it with ``craft.save(doc, name)``."""
    from docx import Document

    return Document(str(files.path(name)))


def replace_in_word(name: str | Path, old: str, new: str, *, save_as: str | None = None) -> int:
    """Replace text in paragraphs, table cells, headers and footers; returns the count.

    A match split across runs keeps the formatting of the run it starts in. Saves
    over the file unless ``save_as`` names a new one.
    """
    if not old:
        raise ValueError("Give the text to replace")
    document = open_word(name)
    count = sum(_replace_in_paragraph(p, old, new) for p in _all_paragraphs(document))
    files.save(document, save_as or name)
    return count


def _all_paragraphs(document: Any) -> list[Any]:
    def from_part(part: Any) -> list[Any]:
        found = list(part.paragraphs)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    found.extend(from_part(cell))
        return found

    paragraphs = from_part(document)
    for section in document.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                paragraphs.extend(from_part(part))
    return paragraphs


def _replace_in_paragraph(paragraph: Any, old: str, new: str) -> int:
    runs = list(paragraph.runs)
    count = 0
    while True:
        texts = [run.text for run in runs]
        joined = "".join(texts)
        at = joined.find(old)
        if at < 0:
            return count
        count += 1
        # Find the runs the match covers; the first keeps the new text, the rest lose theirs.
        start = end = None
        offset = tail = position = 0
        for index, text in enumerate(texts):
            if start is None and at < position + len(text):
                start, offset = index, at - position
            if at + len(old) <= position + len(text):
                end, tail = index, at + len(old) - position
                break
            position += len(text)
        assert start is not None and end is not None
        if start == end:
            runs[start].text = texts[start][:offset] + new + texts[start][tail:]
        else:
            runs[start].text = texts[start][:offset] + new
            for index in range(start + 1, end):
                runs[index].text = ""
            runs[end].text = texts[end][tail:]


# Templates


def template_fields(template: str | Path) -> list[str]:
    """The fields a template fills: Word ``{{ name }}`` fields, Excel defined names, or
    PowerPoint ``{{ name }}`` text in shapes."""
    import re

    source = files.path(template)
    suffix = source.suffix.lower()
    if suffix == ".docx":
        from docxtpl import DocxTemplate

        return sorted(DocxTemplate(str(source)).get_undeclared_template_variables())
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        return sorted(load_workbook(source).defined_names)
    if suffix == ".pptx":
        from pptx import Presentation

        found: set[str] = set()
        for slide in Presentation(str(source)).slides:
            shape: Any
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    found.update(
                        re.findall(r"{{\s*([A-Za-z_][A-Za-z0-9_]*)\s*}}", shape.text_frame.text)
                    )
        return sorted(found)
    raise ValueError("Templates are .docx, .xlsx or .pptx")


def fill_template(template: str | Path, data: dict[str, Any], *, save_as: str | Path) -> str:
    """Fill a template into a new file and return its name; the template is unchanged.

    Word uses ``docxtpl`` (fields, loops and conditions keep the template's styles);
    Excel writes each value to the defined name's cell; PowerPoint replaces
    ``{{ name }}`` in text. A field left unfilled is an error.
    """
    source = files.path(template)
    target = files.path(save_as)
    suffix = source.suffix.lower()
    if target.suffix.lower() != suffix:
        raise ValueError("The filled file keeps the template's extension")
    if target.exists() and target.resolve() == source.resolve():
        raise ValueError("Save the filled document under a new name")
    missing = [field for field in template_fields(template) if field not in data]
    if missing:
        raise ValueError("Missing fields: " + ", ".join(missing))
    target.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".docx":
        from docxtpl import DocxTemplate

        document = DocxTemplate(str(source))
        document.render(data, autoescape=True)
        document.save(str(target))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(source)
        for name, value in data.items():
            defined = book.defined_names.get(name)
            if defined is None:
                continue
            for sheet, reference in defined.destinations:
                cell: Any = book[sheet][reference.replace("$", "").split(":")[0]]
                cell.value = value
        book.save(target)
    else:
        from pptx import Presentation

        deck = Presentation(str(source))
        for slide in deck.slides:
            item: Any
            for item in slide.shapes:
                if getattr(item, "has_text_frame", False):
                    for paragraph in item.text_frame.paragraphs:
                        for name, value in data.items():
                            for spelling in (f"{{{{ {name} }}}}", f"{{{{{name}}}}}"):
                                _replace_in_paragraph(paragraph, spelling, str(value))
        deck.save(str(target))
    return target.relative_to(_root.USER_ROOT).as_posix()


# Building


def report(name: str | Path, title: str, *blocks: Any, subtitle: str | None = None) -> str:
    """Build a Word report from blocks and return its relative name.

    A block is text (``"# "`` to ``"### "`` make headings, ``"- "`` a bullet), a
    DataFrame (a table), a Matplotlib figure or an image file name.
    """
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading(title, level=0)
    if subtitle:
        document.add_paragraph().add_run(subtitle).italic = True
    for block in blocks:
        if isinstance(block, str):
            stripped = block.lstrip()
            level = len(stripped) - len(stripped.lstrip("#"))
            if 1 <= level <= 3 and stripped[level : level + 1] == " ":
                document.add_heading(stripped[level + 1 :], level=level)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(block)
        elif _is_frame(block):
            _word_table(document, block)
        else:
            document.add_picture(_image(block), width=Inches(6))
    return files.save(document, name)


def slides(name: str | Path, title: str, *pages: tuple[str, Any], subtitle: str = "") -> str:
    """Build a PowerPoint deck; each page is ``(title, content)``.

    Content is text (one bullet per line), a DataFrame (a table of up to 15
    rows) or a Matplotlib figure or image file name.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover: Any = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = subtitle
    for heading, content in pages:
        if isinstance(content, str):
            page: Any = deck.slides.add_slide(deck.slide_layouts[1])
            page.shapes.title.text = heading
            page.placeholders[1].text = content
            continue
        slide: Any = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = heading
        if _is_frame(content):
            frame = content.head(15)
            shape = slide.shapes.add_table(
                len(frame) + 1,
                len(frame.columns),
                Inches(0.5),
                Inches(1.5),
                Inches(12.3),
                Inches(5),
            )
            for column, label in enumerate(frame.columns):
                shape.table.cell(0, column).text = str(label)
            for row, values in enumerate(frame.itertuples(index=False), start=1):
                for column, value in enumerate(values):
                    cell = shape.table.cell(row, column)
                    cell.text = str(value)
                    cell.text_frame.paragraphs[0].font.size = Pt(12)
        else:
            slide.shapes.add_picture(_image(content), Inches(1.5), Inches(1.4), height=Inches(5.8))
    return files.save(deck, name)


def workbook(
    name: str | Path, *sheets: tuple[str, Any], charts: dict[str, dict] | None = None
) -> str:
    """Build an Excel workbook; each sheet is ``(sheet name, DataFrame)``.

    ``charts`` maps a sheet name to a Vega-Lite spec whose x and y fields are columns
    of that sheet; it becomes a native Excel bar, line or scatter chart beside the table.
    """
    from openpyxl import Workbook

    book: Any = Workbook()
    book.remove(book.active)
    for title, frame in sheets:
        sheet = book.create_sheet(title[:31])
        sheet.append([str(label) for label in frame.columns])
        for values in frame.itertuples(index=False):
            sheet.append([_cell(value) for value in values])
        spec = (charts or {}).get(title)
        if spec is not None:
            sheet.add_chart(_excel_chart(spec, sheet, list(frame.columns), len(frame)), "H2")
    target = files.path(name)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Workbooks save as .xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target.relative_to(_root.USER_ROOT).as_posix()


def chart_slide(name: str | Path, spec: dict[str, Any], *, title: str | None = None) -> str:
    """Add a native PowerPoint chart from a Vega-Lite spec (inline values, x and y fields)
    as a new slide of an existing deck, or of a new one; returns the deck's name."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    target = files.path(name)
    deck = Presentation(str(target)) if target.exists() else Presentation()
    x, y, rows, mark = _spec_series(spec)
    slide: Any = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = title or _spec_title(spec) or f"{y} by {x}"
    if mark == "point":
        data: Any = XyChartData()
        series = data.add_series(y)
        for row in rows:
            series.add_data_point(row[x], row[y])
        kind = XL_CHART_TYPE.XY_SCATTER
    else:
        data = CategoryChartData()
        data.categories = [str(row[x]) for row in rows]
        data.add_series(y, [row[y] for row in rows])
        kind = XL_CHART_TYPE.LINE_MARKERS if mark == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    slide.shapes.add_chart(kind, Inches(0.7), Inches(1.5), Inches(8.6), Inches(5.2), data)
    return files.save(deck, target)


class Page:
    """One rendered page; shows as an image in a notebook."""

    def __init__(self, number: int, png: bytes):
        self.number, self.png = number, png

    def _repr_png_(self) -> bytes:
        return self.png

    def __repr__(self) -> str:
        return f"<page {self.number}, {len(self.png)} bytes PNG>"


def to_pdf(name: str | Path, *, timeout: float = 120) -> bytes:
    """A Word, Excel or PowerPoint file as PDF bytes, rendered by LibreOffice offline."""
    import subprocess
    import tempfile

    source = files.path(name)
    if source.suffix.lower() not in {".docx", ".xlsx", ".pptx", ".doc", ".xls", ".ppt"}:
        raise ValueError("Pages come from .docx, .xlsx or .pptx files")
    with tempfile.TemporaryDirectory() as out:
        # A persistent profile keeps later conversions warm.
        profile = Path(tempfile.gettempdir(), "craft-office-profile").as_uri()
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={profile}",
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                out,
                str(source),
            ],
            check=True,
            capture_output=True,
            timeout=timeout,
        )
        return Path(out, source.stem + ".pdf").read_bytes()


def pages(name: str | Path, *, first: int = 1, last: int | None = None, dpi: int = 80) -> list:
    """Render pages of an Office or PDF file as PNG images to look at.

    ``display(page)`` shows one in the notebook, and a run_code result passes up to
    two shown images back to the agent.
    """
    import subprocess
    import tempfile

    source = files.path(name)
    pdf = source.read_bytes() if source.suffix.lower() == ".pdf" else to_pdf(name)
    with tempfile.TemporaryDirectory() as out:
        Path(out, "in.pdf").write_bytes(pdf)
        command = ["pdftoppm", "-png", "-r", str(dpi), "-f", str(first)]
        if last is not None:
            command += ["-l", str(last)]
        subprocess.run(
            command + [str(Path(out, "in.pdf")), str(Path(out, "page"))],
            check=True,
            capture_output=True,
            timeout=120,
        )
        images = sorted(Path(out).glob("page-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
        return [Page(int(image.stem.rsplit("-", 1)[1]), image.read_bytes()) for image in images]


def _spec_series(spec: dict[str, Any]) -> tuple[str, str, list[dict[str, Any]], str]:
    encoding = spec.get("encoding", {})
    x, y = encoding.get("x", {}).get("field"), encoding.get("y", {}).get("field")
    values = spec.get("data", {}).get("values")
    if not x or not y or not isinstance(values, list):
        raise ValueError("The spec needs inline data.values and x and y fields")
    mark = spec.get("mark", "bar")
    mark = mark.get("type", "bar") if isinstance(mark, dict) else mark
    return x, y, values, {"line": "line", "point": "point", "circle": "point"}.get(mark, "bar")


def _spec_title(spec: dict[str, Any]) -> str | None:
    title = spec.get("title")
    return title.get("text") if isinstance(title, dict) else title


def _excel_chart(spec: dict[str, Any], sheet: Any, columns: list[Any], rows: int) -> Any:
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart
    from openpyxl.chart.series_factory import SeriesFactory as Series

    x, y, _, mark = _spec_series(spec)
    labels = [str(label) for label in columns]
    if x not in labels or y not in labels:
        raise ValueError("Chart fields must be columns of the sheet")
    x_ref = Reference(sheet, min_col=labels.index(x) + 1, min_row=2, max_row=rows + 1)
    y_ref = Reference(sheet, min_col=labels.index(y) + 1, min_row=1, max_row=rows + 1)
    if mark == "point":
        chart: Any = ScatterChart()
        chart.series.append(Series(y_ref, x_ref, title_from_data=True))
    else:
        chart = LineChart() if mark == "line" else BarChart()
        chart.add_data(y_ref, titles_from_data=True)
        chart.set_categories(x_ref)
    chart.title = _spec_title(spec) or f"{y} by {x}"
    chart.x_axis.title, chart.y_axis.title = x, y
    return chart


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if hasattr(value, "item"):
        return value.item()
    return str(value)


def _is_frame(value: Any) -> bool:
    return type(value).__module__.startswith("pandas") and hasattr(value, "columns")


def _word_table(document: Any, frame: Any) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(label)
    for values in frame.itertuples(index=False):
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = "" if value is None else str(value)


def _image(value: Any) -> Any:
    """An image file path, or a figure rendered to PNG in memory."""
    if isinstance(value, str | Path):
        return str(files.path(value))
    import io

    figure = getattr(value, "figure", None) or value
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", bbox_inches="tight", dpi=150)
    buffer.seek(0)
    return buffer


def _unique(labels: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for label in labels:
        seen[label] = seen.get(label, 0) + 1
        result.append(label if seen[label] == 1 else f"{label}_{seen[label]}")
    return result
