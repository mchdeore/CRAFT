"""Focused generated-data checks of SDK, owned paths and server bindings."""

import asyncio
from types import SimpleNamespace
from uuid import UUID

import pytest
from mcp.types import CallToolResult, ListToolsResult, TextContent


def test_sdk_forwards_existing_names_and_arguments_once():
    from craft import Craft

    class Session:
        def __init__(self):
            self.requests = []

        async def list_tools(self):
            return ListToolsResult(tools=[])

        async def call_tool(self, name, arguments=None, *, meta=None):
            self.requests.append((name, arguments, meta))
            return CallToolResult(
                content=[TextContent(type="text", text='{"hits":[]}')],
                structuredContent={"hits": []},
            )

    async def check():
        session = Session()
        sdk = Craft(session)
        result = await sdk.search("generated source", limit=2)
        assert result.structuredContent == {"hits": []}
        assert len(session.requests) == 1
        name, arguments, meta = session.requests[0]
        assert (name, arguments) == ("search", {"query": "generated source", "limit": 2})
        assert isinstance(meta["craft/toolCallId"], str)
        await sdk.document("itm_012345abcdef", version=1)
        await sdk.skills(limit=10)
        assert [(name, args) for name, args, _ in session.requests[1:]] == [
            ("get_document", {"item": "itm_012345abcdef", "version": 1}),
            ("list_skills", {"limit": 10}),
        ]

        class LostResponse(Session):
            async def call_tool(self, name, arguments=None, *, meta=None):
                self.requests.append((name, arguments, meta))
                raise ConnectionError("response lost")

        lost = LostResponse()
        with pytest.raises(ConnectionError):
            await Craft(lost)._call("make_chart", {"spec": {}, "format": "png"})
        assert len(lost.requests) == 1

    asyncio.run(check())


def test_owned_path_rejects_escape_and_symlink(tmp_path):
    from craft import owned_path

    root = tmp_path / "owner"
    root.mkdir()
    (root / "scripts").mkdir()
    assert owned_path("scripts/a.py", root=root) == root / "scripts/a.py"
    outside = tmp_path / "outside"
    outside.mkdir()
    (root / "link").symlink_to(outside, target_is_directory=True)
    (root / "dangling").symlink_to(outside / "missing")
    for path in ["../outside/a.py", "/etc/passwd", "link/a.py", "dangling", "a/../b", ""]:
        with pytest.raises(ValueError):
            owned_path(path, root=root)


def test_factory_selects_trusted_owner_and_hides_credentials():
    from app.agent.workspace import WorkspaceEndpoint, WorkspaceToolsets
    from app.infrastructure.access import Caller

    owner_a = UUID("00000000-0000-4000-8000-000000000001")
    owner_b = UUID("00000000-0000-4000-8000-000000000002")
    endpoints = {
        str(owner_a): WorkspaceEndpoint(
            url="http://127.0.0.1:8101/mcp",
            headers={"Authorization": "Bearer synthetic-a"},
            tools=("read_file", "execute_code"),
        ),
        str(owner_b): WorkspaceEndpoint(
            url="http://127.0.0.1:8102/mcp",
            headers={"Authorization": "Bearer synthetic-b"},
            tools=("read_file",),
        ),
    }
    factory = WorkspaceToolsets(endpoints)
    with pytest.raises(PermissionError):
        factory.for_caller(Caller(user_id=UUID("00000000-0000-4000-8000-000000000003")))
    selected = factory.for_caller(Caller(user_id=owner_a))
    other = factory.for_caller(Caller(user_id=owner_b))
    try:
        assert selected.url == "http://127.0.0.1:8101/mcp"
        assert other.url == "http://127.0.0.1:8102/mcp"
        provider = selected._header_provider
        assert provider is not None
        assert provider({"url": other.url}) == {"Authorization": "Bearer synthetic-a"}
        assert "synthetic-a" not in repr(endpoints[str(owner_a)])
    finally:
        asyncio.run(selected.close())
        asyncio.run(other.close())


def test_official_mcp_connection_forwards_bearer_outside_arguments(monkeypatch):
    import json

    import httpx
    from craft import client as sdk_module
    from craft.client import connect

    actual_client = httpx.AsyncClient
    requests = []

    def respond(request):
        body = json.loads(request.content)
        requests.append((request, body))
        method = body["method"]
        if "id" not in body:
            return httpx.Response(202)
        if method == "initialize":
            result = {
                "protocolVersion": "2025-06-18",
                "capabilities": {"tools": {}},
                "serverInfo": {"name": "generated-test", "version": "1"},
            }
        elif method == "tools/list":
            result = {
                "tools": [
                    {
                        "name": "search",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "query": {"type": "string"},
                                "limit": {"type": "integer"},
                            },
                        },
                    }
                ]
            }
        elif method == "tools/call":
            result = {"content": [], "structuredContent": {"hits": []}, "isError": False}
        else:
            raise AssertionError(f"Unexpected fake request: {method}")
        return httpx.Response(200, json={"jsonrpc": "2.0", "id": body["id"], "result": result})

    def configured_client(**kwargs):
        return actual_client(transport=httpx.MockTransport(respond), **kwargs)

    monkeypatch.setattr(sdk_module, "httpx", SimpleNamespace(AsyncClient=configured_client))

    async def check():
        async with connect("http://127.0.0.1:9999/mcp", run_token="synthetic-owned-run") as sdk:
            result = await sdk.search("generated evidence", limit=1)
            assert result.structuredContent == {"hits": []}

    asyncio.run(check())
    calls = [(request, body) for request, body in requests if body["method"] == "tools/call"]
    assert len(calls) == 1
    request, body = calls[0]
    assert request.headers["authorization"] == "Bearer synthetic-owned-run"
    assert "cookie" not in request.headers
    assert body["params"]["arguments"] == {"query": "generated evidence", "limit": 1}
    assert "synthetic-owned-run" not in json.dumps(body)


def test_word_edit_roundtrip_preserves_original(tmp_path):
    from shutil import copyfile

    from craft import owned_path
    from docx import Document

    workspace = tmp_path / "owner"
    workspace.mkdir()
    original = owned_path("imports/source.docx", root=workspace)
    output = owned_path("outputs/edited.docx", root=workspace)
    original.parent.mkdir(parents=True)
    generated = Document()
    generated.add_paragraph("Generated source paragraph.")
    generated.add_paragraph("Keep this paragraph.")
    table = generated.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows, [("Metric", "Value"), ("Revenue", "125"), ("Cost", "50")], strict=True
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    generated.save(str(original))
    original_bytes = original.read_bytes()

    output.parent.mkdir(parents=True)
    copyfile(original, output)
    edited = Document(str(output))
    assert [paragraph.text for paragraph in edited.paragraphs] == [
        "Generated source paragraph.",
        "Keep this paragraph.",
    ]
    assert [[cell.text for cell in row.cells] for row in edited.tables[0].rows] == [
        ["Metric", "Value"],
        ["Revenue", "125"],
        ["Cost", "50"],
    ]
    edited.paragraphs[0].text = "Edited paragraph in personal copy."
    edited.tables[0].cell(1, 1).text = "150"
    edited.add_paragraph("Summary: revenue increased by 25.")
    edited.save(str(output))

    saved = Document(str(output))
    assert [paragraph.text for paragraph in saved.paragraphs] == [
        "Edited paragraph in personal copy.",
        "Keep this paragraph.",
        "Summary: revenue increased by 25.",
    ]
    assert [[cell.text for cell in row.cells] for row in saved.tables[0].rows] == [
        ["Metric", "Value"],
        ["Revenue", "150"],
        ["Cost", "50"],
    ]
    assert original.read_bytes() == original_bytes
