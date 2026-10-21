"""Read, edit and build Word, PowerPoint, PDF and Excel files in the workspace.

Names are relative to ``USER_ROOT``. Source documents from the CRAFT library
are read-only; import a personal copy before editing it.
"""

from pathlib import Path
from typing import Any

from craft import files

# Reading


def read_word(name: str | Path) -> str:
    """A Word document as Markdown: headings, paragraphs, lists and tables in order."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(files.path(name)))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                lines.append("#" * int(style[-1]) + " " + text)
            elif style == "Title":
                lines.append("# " + text)
            elif "List" in style:
                lines.append("- " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            rows = [
                [cell.text.strip().replace("|", "\\|") for cell in row.cells]
                for row in Table(child, document).rows
            ]
            if rows:
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("|" + "---|" * len(rows[0]))
                lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n\n".join(lines) + "\n"


def word_tables(name: str | Path) -> list:
    """Every Word table as a DataFrame, using its first row as the header."""
    import pandas as pd
    from docx import Document

    frames = []
    for table in Document(str(files.path(name))).tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            frames.append(pd.DataFrame(rows[1:], columns=pd.Index(_unique(rows[0]))))
    return frames


def read_slides(name: str | Path) -> list[dict[str, Any]]:
    """Each slide's number, title, text and speaker notes."""
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(files.path(name))).slides, start=1):
        heading: Any = slide.shapes.title
        title = heading.text if heading is not None else ""
        shapes: list[Any] = list(slide.shapes)
        texts = [
            shape.text_frame.text
            for shape in shapes
            if shape.has_text_frame and shape != heading and shape.text_frame.text
        ]
        notes_frame: Any = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = notes_frame.text if notes_frame is not None else ""
        slides.append({"slide": number, "title": title, "text": "\n".join(texts), "notes": notes})
    return slides


def read_pdf(name: str | Path) -> list[str]:
    """The text of each page. Scanned pages without a text layer come back empty."""
    from pypdf import PdfReader

    return [page.extract_text() or "" for page in PdfReader(str(files.path(name))).pages]


def sheets(name: str | Path) -> dict:
    """Every sheet of a workbook as a DataFrame, by sheet name. Formulas are not run."""
    return files.read(name, sheet=None)


# Editing


def open_word(name: str | Path):
    """A python-docx Document to edit; save it with ``craft.save(doc, name)``."""
    from docx import Document

    return Document(str(files.path(name)))


def replace_in_word(name: str | Path, old: str, new: str, *, save_as: str | None = None) -> int:
    """Replace text inside paragraph runs and table cells; returns the count.

    Text split across differently formatted runs is not matched. Saves over the
    file unless ``save_as`` names a new one.
    """
    document = open_word(name)
    count = 0
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    files.save(document, save_as or name)
    return count


# Building


def report(name: str | Path, title: str, *blocks: Any, subtitle: str | None = None) -> str:
    """Build a Word report from blocks and return its relative name.

    A block is text (``"# "`` to ``"### "`` make headings, ``"- "`` a bullet), a
    DataFrame (a table), a Matplotlib figure or an image file name.
    """
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading(title, level=0)
    if subtitle:
        document.add_paragraph().add_run(subtitle).italic = True
    for block in blocks:
        if isinstance(block, str):
            stripped = block.lstrip()
            level = len(stripped) - len(stripped.lstrip("#"))
            if 1 <= level <= 3 and stripped[level : level + 1] == " ":
                document.add_heading(stripped[level + 1 :], level=level)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(block)
        elif _is_frame(block):
            _word_table(document, block)
        else:
            document.add_picture(_image(block), width=Inches(6))
    return files.save(document, name)


def slides(name: str | Path, title: str, *pages: tuple[str, Any], subtitle: str = "") -> str:
    """Build a PowerPoint deck; each page is ``(title, content)``.

    Content is text (one bullet per line), a DataFrame (a table of up to 15
    rows) or a Matplotlib figure or image file name.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover: Any = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    cover.placeholders[1].text = subtitle
    for heading, content in pages:
        if isinstance(content, str):
            page: Any = deck.slides.add_slide(deck.slide_layouts[1])
            page.shapes.title.text = heading
            page.placeholders[1].text = content
            continue
        slide: Any = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = heading
        if _is_frame(content):
            frame = content.head(15)
            shape = slide.shapes.add_table(
                len(frame) + 1,
                len(frame.columns),
                Inches(0.5),
                Inches(1.5),
                Inches(12.3),
                Inches(5),
            )
            for column, label in enumerate(frame.columns):
                shape.table.cell(0, column).text = str(label)
            for row, values in enumerate(frame.itertuples(index=False), start=1):
                for column, value in enumerate(values):
                    cell = shape.table.cell(row, column)
                    cell.text = str(value)
                    cell.text_frame.paragraphs[0].font.size = Pt(12)
        else:
            slide.shapes.add_picture(_image(content), Inches(1.5), Inches(1.4), height=Inches(5.8))
    return files.save(deck, name)


def _is_frame(value: Any) -> bool:
    return type(value).__module__.startswith("pandas") and hasattr(value, "columns")


def _word_table(document: Any, frame: Any) -> None:
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(label)
    for values in frame.itertuples(index=False):
        for cell, value in zip(table.add_row().cells, values, strict=True):
            cell.text = "" if value is None else str(value)


def _image(value: Any) -> Any:
    """An image file path, or a figure rendered to PNG in memory."""
    if isinstance(value, str | Path):
        return str(files.path(value))
    import io

    figure = getattr(value, "figure", None) or value
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", bbox_inches="tight", dpi=150)
    buffer.seek(0)
    return buffer


def _unique(labels: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for label in labels:
        seen[label] = seen.get(label, 0) + 1
        result.append(label if seen[label] == 1 else f"{label}_{seen[label]}")
    return result
